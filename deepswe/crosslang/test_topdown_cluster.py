#!/usr/bin/env python3
"""topdown_cluster.py 的语言过滤回归测试。"""

import contextlib
import io
import json
import pathlib
import sys
import tempfile
import unittest
from unittest import mock

import openpyxl

import topdown_cluster as tc


class LanguageFilterTest(unittest.TestCase):
    def _write_trial(self, out_dir, name, language, cycles):
        trial_dir = out_dir / name
        topdown_dir = trial_dir / "topdown"
        topdown_dir.mkdir(parents=True)
        (trial_dir / "meta.json").write_text(
            json.dumps({"language": language}), encoding="utf-8")
        (topdown_dir / tc.STEPS_JSON_NAME).write_text(json.dumps({
            "steps": [{
                "step": 1,
                "n_cmds": 1,
                "wall_s": 1.0,
                "sleep_s": 0.0,
                "counts": {"cpu_cycles": cycles},
                "topdown": {
                    "Retiring": 0.4,
                    "BadSpec": 0.1,
                    "FrontendBound": 0.2,
                    "BackendBound": 0.3,
                },
                "commands": ["echo test"],
            }],
        }), encoding="utf-8")

    def test_parse_languages_accepts_repeated_and_comma_separated_values(self):
        self.assertEqual(
            tc.parse_languages(["Python, javascript", "RUST", "python"]),
            {"python", "javascript", "rust"},
        )

    def test_exclude_lang_rejects_language_outside_trial_set(self):
        argv = [
            "topdown_cluster.py", "unused",
            "--exclude-lang", "python,java",
        ]
        stderr = io.StringIO()
        with mock.patch.object(sys, "argv", argv), \
                contextlib.redirect_stderr(stderr), \
                self.assertRaises(SystemExit) as raised:
            tc.main()

        self.assertEqual(raised.exception.code, 2)
        self.assertIn("不支持的语言: java", stderr.getvalue())
        self.assertIn("python, go, rust, typescript, javascript", stderr.getvalue())

    def test_excluded_language_is_absent_from_all_outputs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = pathlib.Path(tmp)
            out_dir = root / "topdown_out"
            out_dir.mkdir()
            self._write_trial(out_dir, "python-case", "python", 100)
            self._write_trial(out_dir, "go-case", "go", 200)
            json_out = root / "result.json"
            xlsx_dir = root / "clusters"

            argv = [
                "topdown_cluster.py", str(out_dir),
                "--level", "all",
                "--exclude-lang", "PYTHON",
                "--json-out", str(json_out),
                "--xlsx-dir", str(xlsx_dir),
            ]
            with mock.patch.object(sys, "argv", argv), \
                    contextlib.redirect_stdout(io.StringIO()):
                self.assertEqual(tc.main(), 0)

            result = json.loads(json_out.read_text(encoding="utf-8"))
            self.assertEqual(result["excluded_languages"], ["python"])
            self.assertEqual(result["lang_counts"], {"go": 1})
            self.assertEqual(result["n_steps"], 1)

            steps_wb = openpyxl.load_workbook(xlsx_dir / "steps.xlsx", read_only=True)
            steps_rows = list(steps_wb.active.iter_rows(values_only=True))
            self.assertEqual(steps_rows[1][1], "go")
            self.assertEqual(len(steps_rows), 2)

            tools_wb = openpyxl.load_workbook(
                xlsx_dir / "tool_summary.xlsx", read_only=True)
            tool_rows = list(tools_wb.active.iter_rows(values_only=True))
            self.assertEqual(tool_rows[1][-1], "go")
            self.assertEqual(len(tool_rows), 2)


class RepresentativeSelectionTest(unittest.TestCase):
    @staticmethod
    def _member(name, vec, wall_s=1.0):
        return {
            "trial": name,
            "lang": "test",
            "step": 1,
            "cycles": 1,
            "wall_s": wall_s,
            "vec": vec,
            "commands": [f"echo {name}"],
        }

    def test_representative_uses_l2_distance_to_weighted_centroid(self):
        # 质心为 (0.25, 0.25, 0.25, 0.25)。A 的 L2 距离小于 B，
        # 但 L∞ 距离大于 B；因此这个样例可以防止实现退回旧的 L∞ 口径。
        points = [
            ("A", (0.45, 0.25, 0.25, 0.05)),
            ("B", (0.40, 0.40, 0.10, 0.10)),
            ("A-opposite", (0.05, 0.25, 0.25, 0.45)),
            ("B-opposite", (0.10, 0.10, 0.40, 0.40)),
        ]
        members = [self._member(name, vec) for name, vec in points]
        cluster = {
            "members": members,
            "vecs": [m["vec"] for m in members],
            "cycles": 4,
            "wall_s": 4.0,
        }

        summary = tc.summarize_cluster(cluster, total_cyc=4, ci=1)

        self.assertEqual(summary["rep_trial"], "A")
        self.assertEqual(summary["rep_step"], 1)

    def test_merged_row_recomputes_l2_representative_from_all_members(self):
        a = self._member("A", (0.45, 0.25, 0.25, 0.05))
        a_opposite = self._member("A-opposite", (0.05, 0.25, 0.25, 0.45))
        # B 子簇故意给更大的 wall_s：旧逻辑会直接继承它的代表 B。
        b = self._member("B", (0.40, 0.40, 0.10, 0.10), wall_s=5.0)
        b_opposite = self._member("B-opposite", (0.10, 0.10, 0.40, 0.40), wall_s=5.0)
        clusters = [
            {"members": [a, a_opposite], "vecs": [a["vec"], a_opposite["vec"]],
             "cycles": 2, "wall_s": 2.0},
            {"members": [b, b_opposite], "vecs": [b["vec"], b_opposite["vec"]],
             "cycles": 2, "wall_s": 10.0},
        ]
        rows = tc.cluster_to_rows(
            clusters, total_cyc=4, scope="per-tool-type", program="test-tool")

        merged = tc.merge_consecutive_program_rows(rows)

        self.assertEqual(len(merged), 1)
        self.assertEqual(merged[0]["rep_trial"], "A")


if __name__ == "__main__":
    unittest.main()
